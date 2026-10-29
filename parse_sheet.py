#!/usr/bin/env python3
"""
Parse the Horizon Walker wiki workbook (.xlsx) into JSON + optimized WebP images.

Usage:  python scripts/parse_sheet.py <input.xlsx> <output_dir>

Output:
  <output_dir>/data/characters.json   one record per character
  <output_dir>/data/pages.json        standalone tabs (Demigod, Trait, Tier List, ...)
  <output_dir>/data/report.txt        parse report (what was / wasn't found)
  <output_dir>/img/*.webp             optimized, content-hashed images

Design notes
- Sections inside a character block are found by their LABELS ("Stigma",
  "Trait to Get", "Skill (Priority to Lv6)", "Grade :") instead of fixed cell
  addresses, so small layout shifts by the sheet owner don't break parsing.
- Anything that can't be assigned to a section lands in `extra`, so no text is lost.
- Images are matched to sections by their anchor cell (row, col).
"""
import hashlib
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

RARITY_TABS = {"EX Units": "EX", "SS Units": "SS", "S Units": "S", "A Units": "A"}
SKIP_TABS = {"units", "gift"} | {t.lower() for t in RARITY_TABS}
MAX_SIDE = {"portrait": 1000, "art": 800, "icon": 220}

# Spelling differences between tabs (normalized form -> canonical normalized form)
ALIASES = {"martrotho": "matrotho", "yui": "yuimatsumoto"}

COL_B, COL_C, COL_F, COL_G = 2, 3, 6, 7


# ----------------------------------------------------------------- utilities
def norm(name):
    s = unicodedata.normalize("NFKD", str(name)).lower()
    s = re.sub(r"[^a-z0-9]", "", s)
    return ALIASES.get(s, s)


def slug(name):
    s = unicodedata.normalize("NFKD", str(name)).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


class Images:
    """Extracts images from a worksheet, optimizes them, and indexes them by anchor."""

    def __init__(self, out_dir):
        self.dir = Path(out_dir) / "img"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.cache = {}
        self.dims = {}

    def save(self, raw, kind):
        key = (hashlib.sha1(raw).hexdigest()[:16], kind)
        if key in self.cache:
            return self.cache[key]
        name = f"{kind}-{key[0]}.webp"
        path = self.dir / name
        if not path.exists():
            try:
                im = Image.open(io.BytesIO(raw))
                im.load()
                im.thumbnail((MAX_SIDE[kind],) * 2, Image.LANCZOS)
                self.dims[name] = im.size
                has_alpha = im.mode in ("RGBA", "LA", "PA") or "transparency" in im.info
                im = im.convert("RGBA" if has_alpha else "RGB")
                im.save(path, "WEBP", quality=80, method=4)
            except Exception as e:  # corrupt / unsupported image: skip, don't crash
                print(f"  ! image skipped ({e})", file=sys.stderr)
                name = None
        self.cache[key] = name
        return name

    def aspect(self, name):
        if name not in self.dims:
            try:
                self.dims[name] = Image.open(self.dir / name).size
            except Exception:
                return 1
        w, h = self.dims[name]
        return round(w / h, 3) if h else 1

    def index(self, ws):
        """-> list of dicts {row, col, raw} for every floating image on the sheet."""
        out = []
        for im in ws._images:
            try:
                a = im.anchor._from
                ext = getattr(im.anchor, "ext", None)  # displayed size in the sheet (EMU -> px)
                dw = round(ext.width / 9525) if ext is not None and getattr(ext, "width", None) else im.width
                dh = round(ext.height / 9525) if ext is not None and getattr(ext, "height", None) else im.height
                out.append({"row": a.row + 1, "col": a.col + 1, "raw": im._data(),
                            "w": im.width, "h": im.height, "dw": dw, "dh": dh})
            except Exception:
                continue
        return out


def pick_image(imgs, rows, cols, store, kind):
    for im in imgs:
        if im["row"] in rows and im["col"] in cols:
            return store.save(im["raw"], kind)
    return None


# ------------------------------------------------------- character blocks
def find_blocks(ws):
    """Character headers are 2-row merged cells spanning F:L holding the name."""
    heads = []
    for m in ws.merged_cells.ranges:
        if m.min_col == COL_F and m.max_col == 12 and m.max_row - m.min_row == 1:
            name = txt(ws.cell(m.min_row, COL_F).value)
            if name:
                heads.append((m.min_row, name))
    heads.sort()
    blocks = []
    for i, (r, name) in enumerate(heads):
        end = heads[i + 1][0] - 1 if i + 1 < len(heads) else ws.max_row
        blocks.append((name, r, end))
    return blocks


def parse_block(ws, name, r0, r1, rarity, imgs, store):
    cells = {}
    for row in ws.iter_rows(min_row=r0, max_row=r1):
        for c in row:
            t = txt(c.value)
            if t is not None:
                cells[(c.row, c.column)] = t
    used = set()

    def take(key):
        used.add(key)
        return cells[key]

    def first_row(col, pred):
        rows = sorted(r for (r, c), t in cells.items() if c == col and pred(t))
        return rows[0] if rows else None

    stigma_row = first_row(COL_F, lambda t: t.lower() == "stigma")
    trait_row = first_row(COL_F, lambda t: t.lower().startswith("trait to get"))
    skill_row = first_row(COL_F, lambda t: t.lower().startswith("skill (priority"))
    for r, k in ((stigma_row, "stigma"), (trait_row, "trait"), (skill_row, "skill")):
        if r is not None:
            used.add((r, COL_F))
    used.add((r0, COL_F))  # the name header

    block_imgs = [im for im in imgs if r0 <= im["row"] <= r1]
    rec = {"id": slug(name), "name": name, "rarity": rarity, "hasGuide": True,
           "title": None, "profile": {}, "story": None, "weapon": None,
           "stigmas": [], "traits": [], "skills": [], "gifts": [], "extra": []}
    rec["traitsWip"] = bool(trait_row and "wip" in cells[(trait_row, COL_F)].lower())

    # --- portrait: image in column B near the top
    rec["portrait"] = pick_image(block_imgs, range(r0, r0 + 9), {COL_B}, store, "portrait")

    # --- profile: "Label :" in col B, value in col C
    grade_row = None
    for (r, c), t in sorted(cells.items()):
        if c == COL_B and re.match(r"^[A-Za-z ]+\s*:$", t):
            key = t.rstrip(":").strip()
            used.add((r, c))
            if (r, COL_C) in cells:
                rec["profile"][key] = take((r, COL_C))
            if key == "Grade":
                grade_row = r
    if grade_row:
        above = [k for k in cells if k[1] == COL_B and r0 <= k[0] < grade_row and k not in used]
        if above:
            rec["title"] = take(max(above))  # e.g. "Little Witch - Delia"
    story_row = first_row(COL_B, lambda t: t.lower().startswith("background story"))
    over_row = first_row(COL_B, lambda t: t.lower() == "overview")
    rec["overview"] = None
    for label_row, field, stop in ((story_row, "story", over_row), (over_row, "overview", None)):
        if not label_row:
            continue
        used.add((label_row, COL_B))
        # a section runs until the other label (only when that label sits below it)
        limit = stop if stop and stop > label_row else r1 + 1
        parts = [take(k) for k in sorted(cells)
                 if k[1] == COL_B and label_row < k[0] < limit and k not in used]
        rec[field] = "\n".join(parts) or None

    # --- weapon: everything between header and Stigma label (cols F..L)
    w_end = (stigma_row or r0 + 12) - 1
    w_keys = sorted(k for k in cells if r0 < k[0] <= w_end and k[1] >= COL_F and k not in used)
    wname = [take(k) for k in w_keys if k[1] == COL_F]
    wtext = [take(k) for k in w_keys if k[1] > COL_F]
    wimg = pick_image(block_imgs, range(r0, w_end + 1), {COL_F}, store, "art")
    if wname or wtext or wimg:
        rec["weapon"] = {"name": " / ".join(wname) or None, "text": "\n\n".join(wtext) or None, "image": wimg}

    # --- demigod (recommended Mights): label in col >= H, icons below it, names further down
    dg_row = None
    for k in sorted(cells):
        if k[1] >= 8 and cells[k].lower() == "demigod":
            dg_row = k[0]
            used.add(k)
            break
    rec["demigod"] = []
    if dg_row:
        stop = skill_row or r1 + 1
        for im in sorted((i for i in block_imgs if i["col"] >= 8 and dg_row < i["row"] < stop),
                         key=lambda i: i["col"]):
            col_txt = sorted(k for k in cells if k[1] == im["col"] and im["row"] < k[0] < stop and k not in used)
            name = re.sub(r"\s+", " ", take(col_txt[0])) if col_txt else None
            note = re.sub(r"\s+", " ", take(col_txt[1])) if len(col_txt) > 1 else None
            rec["demigod"].append({"name": name, "note": note, "icon": store.save(im["raw"], "icon")})

    # --- stigmas: between Stigma label and next section label
    if stigma_row:
        s_end = min([r for r in (trait_row, skill_row) if r and r > stigma_row] + [r1 + 1]) - 1
        keys = sorted(k for k in cells if stigma_row <= k[0] <= s_end and k[1] >= COL_F and k not in used)
        titles = [k for k in keys if k[1] == COL_G and not re.match(r"^(Set\s*\d|Substat)", cells[k], re.I)
                  and not re.search(r"\bSet\s*\d", cells[k])]
        entries = []
        for i, tk in enumerate(titles):
            span_end = titles[i + 1][0] - 1 if i + 1 < len(titles) else s_end
            entry = {"phase": None, "name": None, "sets": [], "substats": None, "image": None}
            t = take(tk)
            m = re.match(r"^([^:]+):\s*(.+)$", t)
            entry["phase"], entry["name"] = (m.group(1).strip(), m.group(2).strip()) if m else (None, t)
            for k in keys:
                if k not in used and tk[0] <= k[0] <= span_end and re.search(r"\bSet\s*\d", cells[k]):
                    entry["sets"].append(take(k))
            entry["image"] = pick_image(block_imgs, range(tk[0], tk[0] + 2), {COL_F}, store, "art")
            entries.append(entry)
        # substat priority: label cell (+ the list in the cell below when split in two)
        for k in keys:
            if k in used or k[1] < 8 or not re.match(r"^Substat", cells[k], re.I) or not entries:
                continue
            text = take(k)
            below = (k[0] + 1, k[1])
            if below in cells and below not in used and re.match(r"^\d", cells[below]):
                text += "\n" + take(below)
            idx = max([i for i, t in enumerate(titles) if t[0] <= k[0]] or [0])
            entries[idx]["substats"] = (entries[idx]["substats"] + "\n" + text) if entries[idx]["substats"] else text
        rec["stigmas"] = entries
        for k in keys:  # leftovers inside the stigma area
            if k not in used and not re.match(r"^\d", cells[k]):
                rec["extra"].append({"where": "stigma", "text": take(k)})

    # --- traits: name in F, effect in G
    if trait_row:
        t_end = (skill_row if skill_row and skill_row > trait_row else r1 + 1) - 1
        for r in range(trait_row + 1, t_end + 1):
            if (r, COL_F) in cells:
                tr = {"name": take((r, COL_F)), "effect": None}
                if (r, COL_G) in cells:
                    tr["effect"] = take((r, COL_G))
                rec["traits"].append(tr)
            for c in range(COL_G + 1, 13):
                if (r, c) in cells and (r, c) not in used:
                    rec["extra"].append({"where": "trait", "text": take((r, c))})

    # --- skills: one column per skill (F..L): icon image, title, priority number
    if skill_row:
        for c in range(COL_F, 13):
            col_keys = sorted(k for k in cells if k[1] == c and k[0] > skill_row and k not in used)
            if not col_keys and not any(im["col"] == c and im["row"] > skill_row for im in block_imgs):
                continue
            title, prio, desc = None, None, []
            for k in col_keys:
                t = cells[k]
                if re.match(r"^\d+(\.\d+)?$", t):
                    prio = take(k)
                elif title is None:
                    title = take(k)
                else:
                    desc.append(take(k))
            icon = pick_image([im for im in block_imgs if im["row"] > skill_row], range(skill_row, r1 + 1), {c}, store, "icon")
            if title or icon:
                rec["skills"].append({"title": title, "priority": prio, "description": "\n".join(desc) or None, "icon": icon})

    # --- anything unassigned is preserved, never silently dropped
    for k in sorted(cells):
        if k not in used:
            rec["extra"].append({"where": f"R{k[0]}C{k[1]}", "text": cells[k]})
    return rec


def parse_block_forms(ws, name, r0, r1, rarity, imgs, store):
    """A block can hold several forms of one unit (e.g. Griselda + Headless Knight): the second
    form starts right after the first form's skill-priority row and has its own Stigma label."""
    f_lbl = lambda r: (txt(ws.cell(r, COL_F).value) or "").lower()
    stigmas = [r for r in range(r0, r1 + 1) if f_lbl(r) == "stigma"]
    if len(stigmas) < 2:
        return [parse_block(ws, name, r0, r1, rarity, imgs, store)]
    num = re.compile(r"^\d+(\.\d+)?$")
    nums = [r for r in range(stigmas[0], stigmas[1])
            if any(num.match(txt(ws.cell(r, c).value) or "") for c in range(COL_F, 13))]
    split = (max(nums) + 1) if nums else stigmas[1] - 8
    first = parse_block(ws, name, r0, split - 1, rarity, imgs, store)
    second = parse_block(ws, name, split, r1, rarity, imgs, store)
    title = second["title"] or ""
    form_name = title.split(" - ")[-1].strip() if " - " in title else f"{name} (form 2)"
    second["name"], second["id"], second["formOf"] = form_name, slug(form_name), first["id"]
    return [first, second]


# ------------------------------------------------------------ other tabs
def parse_units_tab(ws, imgs, store):
    """Roster order, rarity rows and grid icons (name cell, icon one row above)."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            t = txt(c.value)
            if not t or c.column < 2 or len(t) > 30 or c.row < 8:
                continue
            best, dist = None, 99
            for im in imgs:
                if im["row"] == c.row - 1 and abs(im["col"] - c.column) < dist and im["col"] <= c.column + 1:
                    best, dist = im, abs(im["col"] - c.column)
            if best and dist <= 1:
                out[norm(t)] = {"name": t, "icon": store.save(best["raw"], "icon"), "order": (c.row, c.column)}
    return out


def parse_gift_tab(ws, imgs, store):
    gifts = {}
    for r in range(1, ws.max_row + 1):
        name = txt(ws.cell(r, 1).value)
        if not name or r == 1:
            continue
        items = []
        for c in range(2, ws.max_column + 1):
            t = txt(ws.cell(r, c).value)
            if not t:
                continue
            m = re.match(r"^(.*?)\s*\((\d+)\)\s*$", t)
            icon = pick_image(imgs, {r - 1}, {c}, store, "icon")
            items.append({"name": m.group(1) if m else t, "value": int(m.group(2)) if m else None, "icon": icon})
        gifts[norm(name)] = {"label": name, "items": items}
    return gifts


def parse_generic(ws, imgs, store):
    """Faithful positional dump of a tab: cells (with merge spans + hyperlinks) and images."""
    spans = {(m.min_row, m.min_col): (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
             for m in ws.merged_cells.ranges}
    cells = []
    for row in ws.iter_rows():
        for c in row:
            t = txt(c.value)
            href = c.hyperlink.target if c.hyperlink and c.hyperlink.target else None
            if t is None and not href:
                continue
            rs, cs = spans.get((c.row, c.column), (1, 1))
            d = {"r": c.row, "c": c.column, "t": t}
            if rs > 1:
                d["rs"] = rs
            if cs > 1:
                d["cs"] = cs
            if href:
                d["href"] = href
            cells.append(d)
    images = []
    for im in imgs:
        src = store.save(im["raw"], "art")
        if src:
            images.append({"r": im["row"], "c": im["col"], "src": src, "w": im["dw"], "h": im["dh"], "a": store.aspect(src)})
    return {"cells": cells, "images": images}



# ----------------------------------------------- Trait 2 + Stigma workbooks
PCT = re.compile(r"(Crit Rate|Crit DMG|Boost|ATK|Max HP|Accuracy|Evade|AP Recovery|SPD)")


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("%", "").strip())
    except ValueError:
        return None


def parse_traits(path):
    """Stat matrix sheets -> per trait only the stats that are set. 'Rearranged' and hidden copies are skipped."""
    wb = load_workbook(path, data_only=True)
    tabs = []
    for sheet, tid in (("T1 to T3", "t1-t3"), ("T4 to T6", "t4-t6"), ("T7", "t7")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        stats, idx = [], {}
        for c in ws[1]:
            h = txt(c.value)
            if h and c.column > 1:
                idx[c.column] = len(stats)
                stats.append({"n": h.replace("%", "").strip(),
                              "p": "%" in h or bool(PCT.search(h) and "DEF" not in h)})
        rows = []
        for r in range(2, ws.max_row + 1):
            name = txt(ws.cell(r, 1).value)
            if not name:
                continue
            vals = []
            for col, i in idx.items():
                v = num(ws.cell(r, col).value)
                if v:
                    vals.append([i, v])
            rows.append({"n": name, "s": vals})
        tabs.append({"id": tid, "title": sheet, "kind": "stats", "stats": stats, "rows": rows})
    if "T7 Others" in wb.sheetnames:
        ws = wb["T7 Others"]
        rows = [{"n": txt(ws.cell(r, 1).value), "t": txt(ws.cell(r, 2).value) or ""}
                for r in range(2, ws.max_row + 1) if txt(ws.cell(r, 1).value)]
        tabs.append({"id": "t7-others", "title": "T7 Others", "kind": "text", "rows": rows})
    return tabs


def parse_stigmas(path, store):
    """Each stigma = a 2x4 merged name cell; its picture sits 3 columns to the left."""
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        color = ws.title.replace("Stigma", "").strip()
        imgs = {(i["row"], i["col"]): i for i in store.index(ws)}
        names = sorted((m.min_row, m.min_col, txt(ws.cell(m.min_row, m.min_col).value))
                       for m in ws.merged_cells.ranges
                       if m.max_row - m.min_row == 1 and m.max_col - m.min_col == 3
                       and txt(ws.cell(m.min_row, m.min_col).value))
        for r, c, name in names:
            nxt = [x[0] for x in names if x[1] == c and x[0] > r]
            end = (min(nxt) - 1) if nxt else ws.max_row

            def after(label):
                for rr in range(r + 1, end + 1):
                    t = txt(ws.cell(rr, c).value)
                    if t and re.match(label, t, re.I):
                        for r2 in range(rr + 1, end + 1):
                            v = txt(ws.cell(r2, c).value)
                            if v:
                                return v
                return None

            rec = {"id": slug(color + "-" + name), "name": name, "color": color, "image": None,
                   "set2": after(r"^2\s*set"), "set4": after(r"^4\s*set"), "main": []}
            im = imgs.get((r, c - 3))
            if im:
                rec["image"] = store.save(im["raw"], "art")
            for rr in range(r + 1, end + 1):
                t = txt(ws.cell(rr, c + 3).value)
                if t and not re.match(r"^main stats", t, re.I):
                    m2 = re.match(r"^(\w+):\s*(.+)$", t, re.S)
                    rec["main"].append({"pos": m2.group(1), "text": " ".join(m2.group(2).split())} if m2
                                       else {"pos": "", "text": t})
            out.append(rec)
    return out


# ------------------------------------------------------------------ main
def main(xlsx, out_dir, traits=None, stigma=None):
    out = Path(out_dir)
    (out / "data").mkdir(parents=True, exist_ok=True)
    print(f"Loading {xlsx} ...")
    wb = load_workbook(xlsx)  # keeps hyperlinks + images
    store = Images(out)
    report = []

    units = {}
    if "Units" in wb.sheetnames:
        units = parse_units_tab(wb["Units"], store.index(wb["Units"]), store)
    gifts = {}
    if "Gift" in wb.sheetnames:
        gifts = parse_gift_tab(wb["Gift"], store.index(wb["Gift"]), store)

    characters, seen = [], set()
    for tab, rarity in RARITY_TABS.items():
        if tab not in wb.sheetnames:
            report.append(f"MISSING TAB: {tab}")
            continue
        ws = wb[tab]
        imgs = store.index(ws)
        for name, r0, r1 in find_blocks(ws):
            for rec in parse_block_forms(ws, name, r0, r1, rarity, imgs, store):
                key = norm(rec["name"])
                u = units.get(key)
                rec["icon"] = u["icon"] if u else None
                g = gifts.get(key)
                rec["gifts"] = g["items"] if g else []
                if not g:
                    report.append(f"no gift entry: {rec['name']}")
                if not u:
                    report.append(f"not in Units grid: {rec['name']}")
                seen.add(key)
                characters.append(rec)
                print(f"  {rarity:2} {rec['name']:18} stig={len(rec['stigmas'])} traits={len(rec['traits'])} "
                      f"skills={len(rec['skills'])} gifts={len(rec['gifts'])} extra={len(rec['extra'])}")

    # Roster entries without a guide. Names listed below the last guide row (enemies/NPCs) are ignored.
    guide_rows = [u["order"][0] for k, u in units.items() if k in seen]
    last_row = max(guide_rows) if guide_rows else 0
    row_rarity = {10: "EX", 13: "EX", 16: "EX", 19: "EX", 24: "SS", 29: "S", 34: "A"}
    for key, u in units.items():
        if key not in seen and u["order"][0] > last_row:
            report.append(f"ignored (below guide rows): {u['name']}")
        elif key not in seen:
            characters.append({"id": slug(u["name"]), "name": u["name"], "rarity": row_rarity.get(u["order"][0]), "hasGuide": False,
                               "icon": u["icon"], "portrait": None, "title": None, "profile": {}, "story": None,
                               "weapon": None, "stigmas": [], "traits": [], "skills": [], "extra": [], "demigod": [], "overview": None,
                               "gifts": gifts.get(key, {}).get("items", [])})
            report.append(f"roster entry without guide: {u['name']}")

    # keep roster order from the Units tab (rarity group, then position)
    order = {k: v["order"] for k, v in units.items()}
    characters.sort(key=lambda c: order.get(norm(c["name"]), (10**6, 0)))

    pages = {}
    for ws in wb.worksheets:
        low = ws.title.lower()
        if low in SKIP_TABS or "overview" in low:
            continue
        pages[slug(ws.title)] = {"title": ws.title, **parse_generic(ws, store.index(ws), store)}
        print(f"  page {ws.title}: {len(pages[slug(ws.title)]['cells'])} cells, {len(pages[slug(ws.title)]['images'])} images")

    if traits and Path(traits).exists():
        tabs = parse_traits(traits)
        (out / "data" / "traits2.json").write_text(json.dumps({"tabs": tabs}, ensure_ascii=False), "utf-8")
        print("  trait tabs:", {t["title"]: len(t["rows"]) for t in tabs})
    if stigma and Path(stigma).exists():
        sg = parse_stigmas(stigma, store)
        (out / "data" / "stigmas.json").write_text(json.dumps(sg, ensure_ascii=False), "utf-8")
        print(f"  stigmas: {len(sg)}")

    (out / "data" / "characters.json").write_text(json.dumps(characters, ensure_ascii=False, indent=1), "utf-8")
    (out / "data" / "pages.json").write_text(json.dumps(pages, ensure_ascii=False), "utf-8")
    (out / "data" / "report.txt").write_text("\n".join(report) or "clean", "utf-8")
    print(f"\n{len(characters)} characters, {len(pages)} pages, {len(list((out/'img').glob('*.webp')))} images")
    print("Report:", *report, sep="\n  ")


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("out")
    ap.add_argument("--traits", help="Trait 2 workbook (optional)")
    ap.add_argument("--stigma", help="Stigma workbook (optional)")
    a = ap.parse_args()
    main(a.xlsx, a.out, a.traits, a.stigma)
