import json
from pathlib import Path

from openpyxl import load_workbook

from .characters import find_blocks, parse_block_forms
from .common import RARITY_TABS, SKIP_TABS, norm, slug
from .images import Images
from .sheets import parse_generic, parse_gift_tab, parse_stigmas, parse_traits, parse_units_tab

def main(xlsx, out_dir, traits=None, stigma=None):
    out = Path(out_dir)
    (out / "data").mkdir(parents=True, exist_ok=True)
    print(f"Loading {xlsx} ...")
    wb = load_workbook(xlsx) 
    store = Images(out)
    report = []

    units = {}
    if "Units" in wb.sheetnames:
        units = parse_units_tab(wb["Units"], store.index(wb["Units"]), store)
    gifts = {}
    if "Gift" in wb.sheetnames:
        gifts = parse_gift_tab(wb["Gift"], store.index(wb["Gift"]), store)

    characters, seen = [], set()
    for tab, rarity in RARITY_TABS.items():
        if tab not in wb.sheetnames:
            report.append(f"MISSING TAB: {tab}")
            continue
        ws = wb[tab]
        imgs = store.index(ws)
        for name, r0, r1 in find_blocks(ws):
            for rec in parse_block_forms(ws, name, r0, r1, rarity, imgs, store):
                key = norm(rec["name"])
                u = units.get(key)
                rec["icon"] = u["icon"] if u else None
                g = gifts.get(key)
                rec["gifts"] = g["items"] if g else []
                if not g:
                    report.append(f"no gift entry: {rec['name']}")
                if not u:
                    report.append(f"not in Units grid: {rec['name']}")
                seen.add(key)
                characters.append(rec)
                print(f"  {rarity:2} {rec['name']:18} stig={len(rec['stigmas'])} traits={len(rec['traits'])} "
                      f"skills={len(rec['skills'])} gifts={len(rec['gifts'])} extra={len(rec['extra'])}")

    guide_rows = [u["order"][0] for k, u in units.items() if k in seen]
    last_row = max(guide_rows) if guide_rows else 0
    row_rarity = {10: "EX", 13: "EX", 16: "EX", 19: "EX", 24: "SS", 29: "S", 34: "A"}
    for key, u in units.items():
        if key not in seen and u["order"][0] > last_row:
            report.append(f"ignored (below guide rows): {u['name']}")
        elif key not in seen:
            characters.append({"id": slug(u["name"]), "name": u["name"], "rarity": row_rarity.get(u["order"][0]), "hasGuide": False,
                               "icon": u["icon"], "portrait": None, "title": None, "profile": {}, "story": None,
                               "weapon": None, "stigmas": [], "traits": [], "skills": [], "extra": [], "demigod": [], "overview": None,
                               "gifts": gifts.get(key, {}).get("items", [])})
            report.append(f"roster entry without guide: {u['name']}")

    order = {k: v["order"] for k, v in units.items()}
    characters.sort(key=lambda c: order.get(norm(c["name"]), (10**6, 0)))

    pages = {}
    for ws in wb.worksheets:
        low = ws.title.lower()
        if low in SKIP_TABS or "overview" in low:
            continue
        pages[slug(ws.title)] = {"title": ws.title, **parse_generic(ws, store.index(ws), store)}
        print(f"  page {ws.title}: {len(pages[slug(ws.title)]['cells'])} cells, {len(pages[slug(ws.title)]['images'])} images")

    if traits and Path(traits).exists():
        tabs = parse_traits(traits)
        (out / "data" / "traits2.json").write_text(json.dumps({"tabs": tabs}, ensure_ascii=False), "utf-8")
        print("  trait tabs:", {t["title"]: len(t["rows"]) for t in tabs})
    if stigma and Path(stigma).exists():
        sg = parse_stigmas(stigma, store)
        (out / "data" / "stigmas.json").write_text(json.dumps(sg, ensure_ascii=False), "utf-8")
        print(f"  stigmas: {len(sg)}")

    (out / "data" / "characters.json").write_text(json.dumps(characters, ensure_ascii=False, indent=1), "utf-8")
    (out / "data" / "pages.json").write_text(json.dumps(pages, ensure_ascii=False), "utf-8")
    (out / "data" / "report.txt").write_text("\n".join(report) or "clean", "utf-8")
    print(f"\n{len(characters)} characters, {len(pages)} pages, {len(list((out/'img').glob('*.webp')))} images")
    print("Report:", *report, sep="\n  ")

