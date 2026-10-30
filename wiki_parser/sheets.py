import re

from openpyxl import load_workbook

from .common import norm, slug, txt
from .images import pick_image

def parse_units_tab(ws, imgs, store):
    """Roster order, rarity rows and grid icons (name cell, icon one row above)."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            t = txt(c.value)
            if not t or c.column < 2 or len(t) > 30 or c.row < 8:
                continue
            best, dist = None, 99
            for im in imgs:
                if im["row"] == c.row - 1 and abs(im["col"] - c.column) < dist and im["col"] <= c.column + 1:
                    best, dist = im, abs(im["col"] - c.column)
            if best and dist <= 1:
                out[norm(t)] = {"name": t, "icon": store.save(best["raw"], "icon"), "order": (c.row, c.column)}
    return out


def parse_gift_tab(ws, imgs, store):
    gifts = {}
    for r in range(1, ws.max_row + 1):
        name = txt(ws.cell(r, 1).value)
        if not name or r == 1:
            continue
        items = []
        for c in range(2, ws.max_column + 1):
            t = txt(ws.cell(r, c).value)
            if not t:
                continue
            m = re.match(r"^(.*?)\s*\((\d+)\)\s*$", t)
            icon = pick_image(imgs, {r - 1}, {c}, store, "icon")
            items.append({"name": m.group(1) if m else t, "value": int(m.group(2)) if m else None, "icon": icon})
        gifts[norm(name)] = {"label": name, "items": items}
    return gifts


def parse_generic(ws, imgs, store):
    """Faithful positional dump of a tab: cells (with merge spans + hyperlinks) and images."""
    spans = {(m.min_row, m.min_col): (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
             for m in ws.merged_cells.ranges}
    cells = []
    for row in ws.iter_rows():
        for c in row:
            t = txt(c.value)
            href = c.hyperlink.target if c.hyperlink and c.hyperlink.target else None
            if t is None and not href:
                continue
            rs, cs = spans.get((c.row, c.column), (1, 1))
            d = {"r": c.row, "c": c.column, "t": t}
            if rs > 1:
                d["rs"] = rs
            if cs > 1:
                d["cs"] = cs
            if href:
                d["href"] = href
            cells.append(d)
    images = []
    for im in imgs:
        src = store.save(im["raw"], "art")
        if src:
            images.append({"r": im["row"], "c": im["col"], "src": src, "w": im["dw"], "h": im["dh"], "a": store.aspect(src)})
    return {"cells": cells, "images": images}



PCT = re.compile(r"(Crit Rate|Crit DMG|Boost|ATK|Max HP|Accuracy|Evade|AP Recovery|SPD)")


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("%", "").strip())
    except ValueError:
        return None


def parse_traits(path):
    """Stat matrix sheets -> per trait only the stats that are set. 'Rearranged' and hidden copies are skipped."""
    wb = load_workbook(path, data_only=True)
    tabs = []
    for sheet, tid in (("T1 to T3", "t1-t3"), ("T4 to T6", "t4-t6"), ("T7", "t7")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        stats, idx = [], {}
        for c in ws[1]:
            h = txt(c.value)
            if h and c.column > 1:
                idx[c.column] = len(stats)
                stats.append({"n": h.replace("%", "").strip(),
                              "p": "%" in h or bool(PCT.search(h) and "DEF" not in h)})
        rows = []
        for r in range(2, ws.max_row + 1):
            name = txt(ws.cell(r, 1).value)
            if not name:
                continue
            vals = []
            for col, i in idx.items():
                v = num(ws.cell(r, col).value)
                if v:
                    vals.append([i, v])
            rows.append({"n": name, "s": vals})
        tabs.append({"id": tid, "title": sheet, "kind": "stats", "stats": stats, "rows": rows})
    if "T7 Others" in wb.sheetnames:
        ws = wb["T7 Others"]
        rows = [{"n": txt(ws.cell(r, 1).value), "t": txt(ws.cell(r, 2).value) or ""}
                for r in range(2, ws.max_row + 1) if txt(ws.cell(r, 1).value)]
        tabs.append({"id": "t7-others", "title": "T7 Others", "kind": "text", "rows": rows})
    return tabs


def parse_stigmas(path, store):
    """Each stigma = a 2x4 merged name cell; its picture sits 3 columns to the left."""
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        color = ws.title.replace("Stigma", "").strip()
        imgs = {(i["row"], i["col"]): i for i in store.index(ws)}
        names = sorted((m.min_row, m.min_col, txt(ws.cell(m.min_row, m.min_col).value))
                       for m in ws.merged_cells.ranges
                       if m.max_row - m.min_row == 1 and m.max_col - m.min_col == 3
                       and txt(ws.cell(m.min_row, m.min_col).value))
        for r, c, name in names:
            nxt = [x[0] for x in names if x[1] == c and x[0] > r]
            end = (min(nxt) - 1) if nxt else ws.max_row

            def after(label):
                for rr in range(r + 1, end + 1):
                    t = txt(ws.cell(rr, c).value)
                    if t and re.match(label, t, re.I):
                        for r2 in range(rr + 1, end + 1):
                            v = txt(ws.cell(r2, c).value)
                            if v:
                                return v
                return None

            rec = {"id": slug(color + "-" + name), "name": name, "color": color, "image": None,
                   "set2": after(r"^2\s*set"), "set4": after(r"^4\s*set"), "main": []}
            im = imgs.get((r, c - 3))
            if im:
                rec["image"] = store.save(im["raw"], "art")
            for rr in range(r + 1, end + 1):
                t = txt(ws.cell(rr, c + 3).value)
                if t and not re.match(r"^main stats", t, re.I):
                    m2 = re.match(r"^(\w+):\s*(.+)$", t, re.S)
                    rec["main"].append({"pos": m2.group(1), "text": " ".join(m2.group(2).split())} if m2
                                       else {"pos": "", "text": t})
            out.append(rec)
    return out


